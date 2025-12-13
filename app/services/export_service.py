# app/services/export_service.py
import csv
from io import BytesIO, StringIO
from typing import Optional, List, Dict
from datetime import datetime
import logging
from enum import Enum

from sqlalchemy.orm import Session

from app.models.order import Order

logger = logging.getLogger(__name__)


class ExportFormat(Enum):
    EXCEL = "excel"
    PDF = "pdf"
    CSV = "csv"
    TXT = "txt"


class ExportService:
    def __init__(self, db: Session):
        self.db = db

        # Kiểm tra dependencies
        self._check_dependencies()

    def _check_dependencies(self):
        """Check if required libraries are installed"""
        try:
            import openpyxl
            self.has_excel = True
        except ImportError:
            self.has_excel = False
            logger.warning("openpyxl not installed. Excel export will use CSV instead.")

        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            self.has_pdf = True
        except ImportError:
            self.has_pdf = False
            logger.warning("reportlab not installed. PDF export will use simple text instead.")

    def export_orders_report(self,
                             format_type: ExportFormat,
                             orders: List[Order],
                             start_date: Optional[str] = None,
                             end_date: Optional[str] = None,
                             filter_status: Optional[str] = None) -> Dict:
        """
        Export orders report

        Returns:
            Dict with filename and file content
        """
        logger.info(f"Exporting {len(orders)} orders in {format_type.value} format")

        if format_type == ExportFormat.EXCEL:
            if self.has_excel:
                return self._export_orders_to_excel(orders, start_date, end_date)
            else:
                logger.warning("Falling back to CSV for Excel export")
                return self._export_orders_to_csv(orders, start_date, end_date)

        elif format_type == ExportFormat.PDF:
            if self.has_pdf:
                return self._export_orders_to_pdf(orders, start_date, end_date)
            else:
                logger.warning("Falling back to TXT for PDF export")
                return self._export_orders_to_txt(orders, start_date, end_date)

        elif format_type == ExportFormat.CSV:
            return self._export_orders_to_csv(orders, start_date, end_date)

        elif format_type == ExportFormat.TXT:
            return self._export_orders_to_txt(orders, start_date, end_date)

        else:
            raise ValueError(f"Unsupported format: {format_type}")

    def export_statistics_report(self,
                                 format_type: ExportFormat,
                                 statistics: Dict,
                                 revenue_report: Dict,
                                 best_sellers: List[Dict]) -> Dict:
        """
        Export statistics report
        """
        logger.info(f"Exporting statistics report in {format_type.value} format")

        if format_type == ExportFormat.EXCEL:
            if self.has_excel:
                return self._export_stats_to_excel(statistics, revenue_report, best_sellers)
            else:
                return self._export_stats_to_csv(statistics, revenue_report, best_sellers)

        elif format_type == ExportFormat.CSV:
            return self._export_stats_to_csv(statistics, revenue_report, best_sellers)

        elif format_type == ExportFormat.TXT:
            return self._export_stats_to_txt(statistics, revenue_report, best_sellers)

        else:
            raise ValueError(f"Unsupported format for statistics: {format_type}")

    # ==================== PRIVATE METHODS ====================

    def _export_orders_to_excel(self, orders: List[Order], start_date: Optional[str], end_date: Optional[str]) -> Dict:
        """Export orders to Excel"""
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Orders"

        # Styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_alignment = Alignment(horizontal="center", vertical="center")

        # Headers
        headers = [
            "Mã đơn", "Khách hàng", "Email", "Ngày đặt",
            "Trạng thái", "Phương thức TT", "Trạng thái TT",
            "Tổng tiền", "Phí ship", "Thành tiền",
            "Địa chỉ", "SĐT", "Số SP"
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_alignment
            ws.column_dimensions[chr(64 + col)].width = len(header) + 3

        # Data
        for row_idx, order in enumerate(orders, 2):
            customer_name = order.customer.user.fullname if order.customer and order.customer.user else "N/A"
            email = order.customer.user.email if order.customer and order.customer.user else "N/A"

            row_data = [
                order.order_code,
                customer_name,
                email,
                order.created_at.strftime("%d/%m/%Y %H:%M"),
                order.order_status.value,
                order.payment_method.value,
                order.payment_status.value,
                float(order.total_amount),
                float(order.shipping_fee),
                float(order.final_amount),
                order.shipping_address,
                order.shipping_phone,
                len(order.order_items)
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = center_alignment

        # Save to bytes
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_don_hang_{timestamp}.xlsx",
            "content": output.getvalue(),
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }

    def _export_orders_to_pdf(self, orders: List[Order], start_date: Optional[str], end_date: Optional[str]) -> Dict:
        """Export orders to PDF"""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []

        styles = getSampleStyleSheet()

        # Tiêu đề
        title = Paragraph("BÁO CÁO ĐƠN HÀNG", styles['Title'])
        elements.append(title)
        elements.append(Spacer(1, 0.25 * inch))

        # Thông tin báo cáo
        report_info = f"""
        <b>Thời gian:</b> {start_date or 'Tất cả'} đến {end_date or 'Tất cả'}<br/>
        <b>Ngày xuất:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}<br/>
        <b>Tổng đơn:</b> {len(orders)}<br/>
        <b>Tổng doanh thu:</b> {self._format_currency(sum(o.final_amount for o in orders))}<br/>
        """
        info_paragraph = Paragraph(report_info, styles['Normal'])
        elements.append(info_paragraph)
        elements.append(Spacer(1, 0.5 * inch))

        # Bảng dữ liệu
        table_data = []
        headers = ["Mã đơn", "Khách hàng", "Ngày", "Trạng thái", "PTTT", "Thành tiền"]
        table_data.append(headers)

        for order in orders:
            customer_name = order.customer.user.fullname if order.customer and order.customer.user else "N/A"
            customer_name = customer_name[:15] + "..." if len(customer_name) > 15 else customer_name

            table_data.append([
                order.order_code,
                customer_name,
                order.created_at.strftime("%d/%m"),
                order.order_status.value[:10],
                order.payment_method.value[:3],
                self._format_currency(order.final_amount, short=True)
            ])

        table = Table(table_data, colWidths=[1.2 * inch, 1.2 * inch, 0.6 * inch, 0.8 * inch, 0.5 * inch, 0.8 * inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
        ]))

        elements.append(table)
        doc.build(elements)
        buffer.seek(0)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_don_hang_{timestamp}.pdf",
            "content": buffer.getvalue(),
            "content_type": "application/pdf"
        }

    def _export_orders_to_csv(self, orders: List[Order], start_date: Optional[str], end_date: Optional[str]) -> Dict:
        """Export orders to CSV"""
        output = StringIO()
        writer = csv.writer(output)

        # Header (tiếng Việt có dấu)
        headers = [
            "Mã đơn", "Khách hàng", "Email", "Ngày đặt hàng",
            "Trạng thái", "Phương thức thanh toán", "Trạng thái thanh toán",
            "Tổng tiền", "Phí vận chuyển", "Thành tiền",
            "Địa chỉ giao hàng", "Số điện thoại", "Số sản phẩm"
        ]
        writer.writerow(headers)

        # Data
        for order in orders:
            customer_name = order.customer.user.fullname if order.customer and order.customer.user else "N/A"
            email = order.customer.user.email if order.customer and order.customer.user else "N/A"

            writer.writerow([
                order.order_code,
                customer_name,
                email,
                order.created_at.strftime("%d/%m/%Y %H:%M"),
                order.order_status.value,
                order.payment_method.value,
                order.payment_status.value,
                f"{order.total_amount:.2f}",
                f"{order.shipping_fee:.2f}",
                f"{order.final_amount:.2f}",
                order.shipping_address,
                order.shipping_phone,
                len(order.order_items)
            ])

        csv_bytes = output.getvalue().encode('utf-8')
        output.close()

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_don_hang_{timestamp}.csv",
            "content": csv_bytes,
            "content_type": "text/csv; charset=utf-8"
        }

    def _export_orders_to_txt(self, orders: List[Order], start_date: Optional[str], end_date: Optional[str]) -> Dict:
        """Export orders to TXT"""
        output = StringIO()

        output.write("=" * 60 + "\n")
        output.write("BÁO CÁO ĐƠN HÀNG\n")
        output.write("=" * 60 + "\n\n")

        output.write(f"Thời gian: {start_date or 'Tất cả'} đến {end_date or 'Tất cả'}\n")
        output.write(f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}\n")
        output.write(f"Tổng số đơn: {len(orders)}\n")

        total_revenue = sum(order.final_amount for order in orders)
        output.write(f"Tổng doanh thu: {self._format_currency(total_revenue)}\n\n")

        output.write("-" * 60 + "\n")
        output.write("CHI TIẾT ĐƠN HÀNG\n")
        output.write("-" * 60 + "\n\n")

        for i, order in enumerate(orders, 1):
            customer_name = order.customer.user.fullname if order.customer and order.customer.user else "N/A"

            output.write(f"Đơn #{i}: {order.order_code}\n")
            output.write(f"  Khách hàng: {customer_name}\n")
            output.write(f"  Ngày đặt: {order.created_at.strftime('%d/%m/%Y %H:%M')}\n")
            output.write(f"  Trạng thái: {order.order_status.value}\n")
            output.write(f"  Thanh toán: {order.payment_method.value} ({order.payment_status.value})\n")
            output.write(f"  Thành tiền: {self._format_currency(order.final_amount)}\n")
            output.write(f"  Số sản phẩm: {len(order.order_items)}\n")
            output.write(f"  Địa chỉ: {order.shipping_address}\n")
            output.write(f"  SĐT: {order.shipping_phone}\n")
            output.write("-" * 40 + "\n")

        txt_bytes = output.getvalue().encode('utf-8')
        output.close()

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_don_hang_{timestamp}.txt",
            "content": txt_bytes,
            "content_type": "text/plain; charset=utf-8"
        }

    def _export_stats_to_excel(self, statistics: Dict, revenue_report: Dict, best_sellers: List[Dict]) -> Dict:
        """Export statistics to Excel"""
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Thống kê"

        # Tiêu đề
        ws['A1'] = "BÁO CÁO THỐNG KÊ"
        ws['A1'].font = Font(bold=True, size=14)

        # Thông tin chung
        ws['A3'] = "THỐNG KÊ CHUNG"
        ws['A3'].font = Font(bold=True)

        stats_data = [
            ["Tổng số đơn hàng:", statistics["total_orders"]],
            ["Tổng doanh thu:", f"{statistics['total_revenue']:,.0f} VND"],
            ["Đơn hàng hôm nay:", statistics["today_orders"]],
            ["Đơn hàng tháng này:", statistics["this_month_orders"]],
        ]

        for i, (label, value) in enumerate(stats_data, 4):
            ws.cell(row=i, column=1, value=label)
            ws.cell(row=i, column=2, value=value)

        # Đơn hàng theo trạng thái
        row_start = len(stats_data) + 6
        ws.cell(row=row_start, column=1, value="ĐƠN HÀNG THEO TRẠNG THÁI")
        ws.cell(row=row_start, column=1).font = Font(bold=True)

        for i, (status, count) in enumerate(statistics["orders_by_status"].items(), row_start + 1):
            ws.cell(row=i, column=1, value=f"  {status}:")
            ws.cell(row=i, column=2, value=count)

        # Sản phẩm bán chạy
        row_start += len(statistics["orders_by_status"]) + 3
        ws.cell(row=row_start, column=1, value="SẢN PHẨM BÁN CHẠY (Top 10)")
        ws.cell(row=row_start, column=1).font = Font(bold=True)

        headers = ["STT", "Mã SP", "Tên sản phẩm", "Số lượng", "Doanh thu"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row_start + 1, column=col, value=header)
            cell.font = Font(bold=True)

        for i, product in enumerate(best_sellers[:10], row_start + 2):
            ws.cell(row=i, column=1, value=i - row_start - 1)
            ws.cell(row=i, column=2, value=product["product_id"])
            ws.cell(row=i, column=3, value=product["product_name"])
            ws.cell(row=i, column=4, value=product["total_quantity"])
            ws.cell(row=i, column=5, value=f"{product['total_revenue']:,.0f}")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_thong_ke_{timestamp}.xlsx",
            "content": output.getvalue(),
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }

    def _export_stats_to_csv(self, statistics: Dict, revenue_report: Dict, best_sellers: List[Dict]) -> Dict:
        """Export statistics to CSV"""
        output = StringIO()
        writer = csv.writer(output)

        writer.writerow(["BÁO CÁO THỐNG KÊ"])
        writer.writerow([f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        writer.writerow([])

        writer.writerow(["THỐNG KÊ CHUNG"])
        writer.writerow(["Chỉ số", "Giá trị"])
        writer.writerow(["Tổng số đơn hàng", statistics["total_orders"]])
        writer.writerow(["Tổng doanh thu", f"{statistics['total_revenue']:,.0f} VND"])
        writer.writerow(["Đơn hàng hôm nay", statistics["today_orders"]])
        writer.writerow(["Đơn hàng tháng này", statistics["this_month_orders"]])
        writer.writerow([])

        writer.writerow(["ĐƠN HÀNG THEO TRẠNG THÁI"])
        writer.writerow(["Trạng thái", "Số lượng"])
        for status, count in statistics["orders_by_status"].items():
            writer.writerow([status, count])
        writer.writerow([])

        writer.writerow(["SẢN PHẨM BÁN CHẠY"])
        writer.writerow(["STT", "Mã SP", "Tên sản phẩm", "Số lượng", "Doanh thu"])
        for i, product in enumerate(best_sellers[:20], 1):
            writer.writerow([
                i,
                product["product_id"],
                product["product_name"],
                product["total_quantity"],
                f"{product['total_revenue']:,.0f}"
            ])

        csv_bytes = output.getvalue().encode('utf-8')
        output.close()

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_thong_ke_{timestamp}.csv",
            "content": csv_bytes,
            "content_type": "text/csv; charset=utf-8"
        }

    def _export_stats_to_txt(self, statistics: Dict, revenue_report: Dict, best_sellers: List[Dict]) -> Dict:
        """Export statistics to TXT"""
        output = StringIO()

        output.write("=" * 60 + "\n")
        output.write("BÁO CÁO THỐNG KÊ\n")
        output.write("=" * 60 + "\n\n")

        output.write(f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}\n\n")

        output.write("THỐNG KÊ CHUNG:\n")
        output.write(f"  • Tổng số đơn hàng: {statistics['total_orders']}\n")
        output.write(f"  • Tổng doanh thu: {self._format_currency(statistics['total_revenue'])}\n")
        output.write(f"  • Đơn hàng hôm nay: {statistics['today_orders']}\n")
        output.write(f"  • Đơn hàng tháng này: {statistics['this_month_orders']}\n\n")

        output.write("ĐƠN HÀNG THEO TRẠNG THÁI:\n")
        for status, count in statistics["orders_by_status"].items():
            output.write(f"  • {status}: {count}\n")
        output.write("\n")

        output.write("ĐƠN HÀNG THEO TRẠNG THÁI THANH TOÁN:\n")
        for status, count in statistics["orders_by_payment"].items():
            output.write(f"  • {status}: {count}\n")
        output.write("\n")

        output.write("SẢN PHẨM BÁN CHẠY (Top 10):\n")
        for i, product in enumerate(best_sellers[:10], 1):
            output.write(f"  {i}. {product['product_name']}\n")
            output.write(f"     • Mã SP: {product['product_id']}\n")
            output.write(f"     • Số lượng: {product['total_quantity']}\n")
            output.write(f"     • Doanh thu: {self._format_currency(product['total_revenue'])}\n")

        txt_bytes = output.getvalue().encode('utf-8')
        output.close()

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return {
            "filename": f"bao_cao_thong_ke_{timestamp}.txt",
            "content": txt_bytes,
            "content_type": "text/plain; charset=utf-8"
        }

    def _format_currency(self, amount: float, short: bool = False) -> str:
        """Format currency for display"""
        if short:
            if amount >= 1000000:
                return f"{amount / 1000000:.1f}M VND"
            elif amount >= 1000:
                return f"{amount / 1000:.1f}K VND"

        return f"{amount:,.0f} VND"